from __future__ import annotations

import argparse
import io
import json
import time
from pathlib import Path
from statistics import mean
from typing import Any

from docx import Document
from pypdf import PdfWriter

import document_extractor
from config import BASE_DIR

EVAL_DIR = BASE_DIR / "eval"
DEFAULT_CASES = EVAL_DIR / "document_extraction_cases.json"
DEFAULT_JSON_OUTPUT = EVAL_DIR / "document_extraction_evaluation.json"
DEFAULT_MD_OUTPUT = EVAL_DIR / "document_extraction_evaluation.md"


def load_cases(path: Path) -> list[dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8")).get("cases", [])


def pdf_bytes(pages: list[str]) -> bytes:
    objects: list[bytes] = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{' '.join(f'{3 + index * 3} 0 R' for index in range(len(pages)))}] /Count {len(pages)} >> endobj\n".encode("latin-1"),
    ]
    for index, text in enumerate(pages):
        page_id = 3 + index * 3
        font_id = page_id + 1
        content_id = page_id + 2
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.extend(
            [
                f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >> endobj\n".encode("latin-1"),
                f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode("latin-1"),
                f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode("latin-1") + stream + b"\nendstream endobj\n",
            ]
        )
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for obj in objects:
        offsets.append(len(output))
        output.extend(obj)
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1"))
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("latin-1"))
    return bytes(output)


def blank_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(case: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    for paragraph in case.get("paragraphs", []):
        document.add_paragraph(paragraph)
    table_rows = case.get("table") or []
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=max(len(row) for row in table_rows))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(buffer)
    return buffer.getvalue()


def case_content(case: dict[str, Any]) -> tuple[bytes, str | None]:
    kind = case["kind"]
    if kind == "pdf":
        return pdf_bytes(case.get("pages", ["Synthetic PDF text."])), "application/pdf"
    if kind == "blank_pdf":
        return blank_pdf_bytes(), "application/pdf"
    if kind == "docx":
        return docx_bytes(case), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if kind == "txt":
        return case.get("text", "").encode("utf-8"), "text/plain"
    if kind == "oversized":
        return b"x" * (document_extractor.max_upload_bytes() + 1), case.get("content_type", "text/plain")
    if kind == "byte_values":
        return bytes(case.get("payload_bytes", [])), case.get("content_type")
    return case.get("payload", "").encode("utf-8"), case.get("content_type")


def evaluate_case(case: dict[str, Any]) -> dict[str, Any]:
    content, content_type = case_content(case)
    started = time.perf_counter()
    error = None
    try:
        extraction = document_extractor.extract_document(case["filename"], content, content_type)
        status = extraction.status
        payload = extraction.to_dict()
    except document_extractor.DocumentExtractionError as exc:
        status = "rejected_oversized" if "too large" in str(exc).lower() else "failed"
        payload = {
            "status": status,
            "filename": case["filename"],
            "file_type": None,
            "size_bytes": len(content),
            "page_count": None,
            "character_count": 0,
            "text": "",
            "pages": [],
            "warnings": [str(exc)],
            "processing_time_ms": None,
        }
        error = exc.__class__.__name__
    latency_ms = int((time.perf_counter() - started) * 1000)
    expected_status = set(case.get("expected_status", []))
    expected_text = case.get("expected_text", [])
    expected_page_count = case.get("expected_page_count")
    status_ok = status in expected_status
    text_ok = all(fragment in payload.get("text", "") for fragment in expected_text)
    page_ok = expected_page_count is None or payload.get("page_count") == expected_page_count
    return {
        "id": case["id"],
        "filename": case["filename"],
        "expected_status": sorted(expected_status),
        "actual_status": status,
        "latency_ms": latency_ms,
        "character_count": payload.get("character_count", 0),
        "page_count": payload.get("page_count"),
        "warnings": payload.get("warnings", []),
        "error": error,
        "correct": status_ok and text_ok and page_ok,
    }


def pct(count: int, total: int) -> float:
    return round(count / total, 3) if total else 0.0


def evaluate(path: Path) -> dict[str, Any]:
    results = [evaluate_case(case) for case in load_cases(path)]
    total = len(results)
    latencies = [item["latency_ms"] for item in results]
    ocr_cases = [item for item in results if "ocr_required" in item["expected_status"]]
    unsupported_cases = [item for item in results if "unsupported" in item["expected_status"]]
    oversized_cases = [item for item in results if "rejected_oversized" in item["expected_status"]]
    successes = [item for item in results if item["actual_status"] in {"success", "partial"}]
    return {
        "evaluation_config": {"case_file": path.as_posix(), "gemini_used": False},
        "metrics": {
            "case_count": total,
            "correct_cases": sum(1 for item in results if item["correct"]),
            "overall_accuracy": pct(sum(1 for item in results if item["correct"]), total),
            "successful_extractions": len(successes),
            "extraction_success_rate": pct(len(successes), total),
            "ocr_required_detection_rate": pct(sum(1 for item in ocr_cases if item["actual_status"] == "ocr_required"), len(ocr_cases)),
            "unsupported_rejection_rate": pct(sum(1 for item in unsupported_cases if item["actual_status"] == "unsupported"), len(unsupported_cases)),
            "oversized_rejection_rate": pct(sum(1 for item in oversized_cases if item["actual_status"] == "rejected_oversized"), len(oversized_cases)),
            "processing_failures": sum(1 for item in results if item["actual_status"] == "failed"),
            "average_latency_ms": round(mean(latencies), 1) if latencies else 0,
            "max_latency_ms": max(latencies) if latencies else 0,
        },
        "results": results,
        "failures": [item for item in results if not item["correct"]],
    }


def percent(value: float) -> str:
    return f"{value * 100:.1f}%"


def render_markdown(report: dict[str, Any]) -> str:
    metrics = report["metrics"]
    lines = [
        "# Document Extraction Evaluation",
        "",
        "## Overall Results",
        f"- Cases: {metrics['case_count']}",
        f"- Overall accuracy: {percent(metrics['overall_accuracy'])}",
        f"- Extraction success rate: {percent(metrics['extraction_success_rate'])}",
        f"- OCR-required detection: {percent(metrics['ocr_required_detection_rate'])}",
        f"- Unsupported-file rejection: {percent(metrics['unsupported_rejection_rate'])}",
        f"- Oversized-file rejection: {percent(metrics['oversized_rejection_rate'])}",
        f"- Average latency: {metrics['average_latency_ms']} ms",
        "",
        "## Case Results",
    ]
    for item in report["results"]:
        lines.append(f"- `{item['id']}` -> {item['actual_status']} ({'pass' if item['correct'] else 'review'})")
    lines.extend(["", "## Failures"])
    if not report["failures"]:
        lines.append("- No extraction evaluation failures.")
    for item in report["failures"]:
        lines.append(f"- `{item['id']}` expected {item['expected_status']}, got {item['actual_status']}")
    lines.extend(
        [
            "",
            "## Recommendation",
            "- Extraction is ready for Part 10B structured fact extraction after user confirmation is designed.",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate local user-document text extraction.")
    parser.add_argument("--cases", type=Path, default=DEFAULT_CASES)
    parser.add_argument("--json-output", type=Path, default=DEFAULT_JSON_OUTPUT)
    parser.add_argument("--md-output", type=Path, default=DEFAULT_MD_OUTPUT)
    args = parser.parse_args()
    report = evaluate(args.cases)
    args.json_output.write_text(json.dumps(report, indent=2, sort_keys=True), encoding="utf-8")
    args.md_output.write_text(render_markdown(report), encoding="utf-8")
    metrics = report["metrics"]
    print(
        "Document extraction evaluation complete: "
        f"{metrics['case_count']} cases; "
        f"accuracy={percent(metrics['overall_accuracy'])}; "
        f"success={percent(metrics['extraction_success_rate'])}"
    )


if __name__ == "__main__":
    main()
