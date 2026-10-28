"""Evaluation for the Case File export (case_file.py).

Seeds each fixture case in eval/case_file_cases.json into a throwaway SQLite DB through the real
storage path, builds the case file context, renders a real DOCX, and checks:
  - every expected section is present and every expected-absent section is omitted (in the
    context AND in the rendered DOCX text -- no empty headings);
  - every cited source resolves against the parsed corpus (chunk id exists, its document matches,
    and the excerpt is the start of that chunk's text);
  - unverified / unrecorded sources are excluded; authority labels are right;
  - facts are labelled "as stated by you" (or name their document), corrections replace, and
    sensitive values never reach the file.

No LLM calls; deterministic. Run: python evaluate_case_file.py
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sqlite3
import tempfile
import time
from functools import lru_cache
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from docx import Document

import case_file
import case_store
import document_store
import grounded_answer
from config import BASE_DIR

EVAL_DIR = BASE_DIR / "eval"
CHUNKS_PATH = BASE_DIR / "parsed" / "legal_chunks.jsonl"
DEFAULT_CASE_PATH = EVAL_DIR / "case_file_cases.json"
DEFAULT_JSON_OUTPUT = EVAL_DIR / "case_file_evaluation.json"
DEFAULT_MD_OUTPUT = EVAL_DIR / "case_file_evaluation.md"

SECTION_HEADINGS = {
    "situation": "Your situation in brief",
    "facts": "Facts as stated by you",
    "timeline": "Timeline",
    "laws": "Laws that may apply",
    "where_to_go": "Where to go",
    "evidence": "Evidence checklist",
    "lawyer_questions": "Questions you may want to ask the lawyer",
    "free_legal_help": "Free legal help",
    "appendix": "Appendix — conversation record",
}


def load_cases(path: Path = DEFAULT_CASE_PATH) -> list[dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8")).get("cases", [])


@lru_cache(maxsize=1)
def corpus_chunks() -> dict[str, dict[str, Any]]:
    chunks = {}
    with CHUNKS_PATH.open(encoding="utf-8") as handle:
        for line in handle:
            item = json.loads(line)
            chunks[item["chunk_id"]] = item
    return chunks


def _chunk_object(chunk_id: str) -> SimpleNamespace:
    item = corpus_chunks()[chunk_id]
    return SimpleNamespace(**{**item, "source": item.get("source") or item.get("source_file")})


def seed_case(spec: dict[str, Any], db_path: Path, user_id: int | None = None) -> int:
    """Store one fixture case exactly the way /api/ask would: create_case, save_message with
    assistant_content_for_storage() (so verification is compacted by the real code), and a
    documents row carrying the confirmed-fact context (as /api/documents/confirm-facts does)."""
    document_store.init_db(db_path)
    contexts = {}
    for document in spec.get("documents") or []:
        confirmed_facts = {"document_type": document["document_type"], "document_summary": "User-confirmed facts from document."}
        for category, items in document["facts"].items():
            confirmed_facts[category] = [{**item, "source_page": 1, "source_excerpt": item["value"], "confidence": "high"} for item in items]
        context = {
            "confirmed_fact_context_id": document["context_id"],
            "source": "user_document",
            "status": "confirmed",
            "document_type": document["document_type"],
            "confirmed_facts": confirmed_facts,
        }
        contexts[document["context_id"]] = context
        now = case_store.utc_timestamp()
        with case_store.connect(db_path) as conn:
            conn.execute(
                "INSERT INTO documents (original_filename, display_filename, file_type, mime_type, size_bytes, created_at, "
                "updated_at, storage_path, extraction_status, confirmed_fact_context_json, user_id) VALUES (?, ?, 'pdf', 'application/pdf', 1, ?, ?, '', 'success', ?, ?)",
                (
                    document["filename"],
                    document["filename"],
                    now,
                    now,
                    case_store.json_dumps({"status": "confirmed", "confirmed_fact_context_id": document["context_id"], "confirmed_facts": confirmed_facts}),
                    user_id,
                ),
            )

    state = {"conversation_state_id": f"state_{spec['id']}", "primary_domain": spec["primary_domain"], **spec["conversation_state"]}
    first_context = next((contexts[turn["context_id"]] for turn in spec["turns"] if turn.get("context_id")), None)
    case_id = case_store.create_case(spec["turns"][0]["user"], spec["primary_domain"], conversation_state=state, confirmed_document_context=first_context, db_path=db_path, user_id=user_id)
    case_store.rename_case(case_id, spec["title"], db_path, user_id=user_id)
    for turn in spec["turns"]:
        case_store.save_message(case_id, "user", case_store.user_content_for_storage(turn["user"]), db_path)
        case_store.save_message(case_id, "assistant", case_store.assistant_content_for_storage(_response_for(turn, state, contexts)), db_path)
    return case_id


def _response_for(turn: dict[str, Any], state: dict[str, Any], contexts: dict[str, Any]) -> dict[str, Any]:
    reply = turn["reply"]
    turn_state = {**state, "last_user_intent": turn["turn_type"]}
    if reply["kind"] == "ack":
        return {
            "answer": {"issue_summary": reply["message"], "possible_rights": [], "next_steps": []},
            "sources": [],
            "confidence": "high",
            "insufficient_context": False,
            "conversation": {"turn_type": turn["turn_type"]},
            "conversation_state": turn_state,
        }
    chunk_ids = reply["sources"]
    chunks = {chunk_id: _chunk_object(chunk_id) for chunk_id in chunk_ids}
    response = {
        "answer": {
            "issue_summary": reply["issue_summary"],
            "what_this_may_involve": [],
            "possible_legal_position": ["The cited provision may be relevant."],
            "suggested_next_steps": [],
            "evidence_to_preserve": reply.get("evidence_to_preserve", []),
            "where_to_approach": reply.get("where_to_approach", []),
            "limitations": [],
        },
        "sources": grounded_answer.sources_from_ids(chunk_ids, chunks),
        "confidence": "medium",
        "insufficient_context": False,
        "routing": {"status": "classified", "domains": state["domains"], "primary_domain": state["primary_domain"], "confidence": "high"},
        "conversation_state": turn_state,
    }
    if reply.get("verification") != "not_recorded":
        results = [{"claim_id": "claim_001", "support": "supported", "supporting_chunk_ids": reply.get("verified", [])}]
        if reply.get("unverified"):
            results.append({"claim_id": "claim_002", "support": "unsupported", "supporting_chunk_ids": reply["unverified"]})
        response["verification"] = {"status": "verified", "results": results, "claims": []}
    if turn.get("context_id"):
        context = contexts[turn["context_id"]]
        response["confirmed_fact_context_id"] = context["confirmed_fact_context_id"]
        response["document_evidence_used"] = [{"label": "fact", "value": "placeholder", "document_type": context["document_type"]}]
    return response


def docx_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def docx_headings(docx_bytes: bytes) -> list[str]:
    document = Document(io.BytesIO(docx_bytes))
    return [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]


def source_resolves(law: dict[str, Any]) -> bool:
    chunk = corpus_chunks().get(law.get("chunk_id") or "")
    if not chunk:
        return False
    if (chunk.get("source") or chunk.get("source_file")) != law["source_file"]:
        return False
    excerpt = law["excerpt"].rstrip(".").rstrip()
    normalized_chunk = re.sub(r"\s+", " ", chunk.get("text", "")).strip()
    return bool(excerpt) and normalized_chunk.startswith(excerpt[: max(40, len(excerpt) - 5)])


def evaluate_case(spec: dict[str, Any]) -> dict[str, Any]:
    started = time.perf_counter()
    expected = spec["expected"]
    with tempfile.TemporaryDirectory() as tmp:
        db_path = Path(tmp) / "case_file_eval.db"
        case_id = seed_case(spec, db_path)
        bundle = case_file.load_case_for_export(case_id, db_path)
        context = case_file.build_case_file_context(bundle)
        docx_bytes = case_file.render_case_file_docx(context)
    text = docx_text(docx_bytes)
    headings = docx_headings(docx_bytes)
    all_fact_text = "\n".join(fact["text"] for fact in context["facts"])

    checks: dict[str, bool] = {}
    for section in expected["sections_present"]:
        in_docx = section == "cover" or SECTION_HEADINGS[section] in headings
        checks[f"present:{section}"] = section in context["sections"] and in_docx
    for section in expected["sections_absent"]:
        checks[f"absent:{section}"] = section not in context["sections"] and SECTION_HEADINGS[section] not in headings
    checks["no_empty_headings"] = all(heading.strip() for heading in headings)
    checks["no_deadline_or_forum_finder_section"] = not any("deadline" in h.lower() or "forum finder" in h.lower() for h in headings)
    checks["not_legal_advice_notice_on_cover"] = case_file.constants.NOT_LEGAL_ADVICE_NOTICE in text
    for fragment in expected["facts_include"]:
        checks[f"fact_included:{fragment}"] = fragment in all_fact_text
    for fragment in expected["facts_exclude"]:
        checks[f"fact_excluded:{fragment}"] = fragment not in all_fact_text
    checks["facts_labelled"] = all(fact["label"].startswith(("As stated by you", "From your document", "From a document you confirmed")) for fact in context["facts"])
    if expected.get("document_fact_label_contains"):
        checks["document_facts_name_document"] = any(expected["document_fact_label_contains"] in fact["label"] for fact in context["facts"])
    if expected.get("timeline_order"):
        dates = [event["date_display"] for event in context["timeline"]]
        checks["timeline_sorted_and_complete"] = dates == sorted(set(dates), key=dates.index) and all(d in dates for d in expected["timeline_order"]) and [d for d in dates if d in expected["timeline_order"]] == expected["timeline_order"]
    if expected.get("situation_label_contains"):
        checks["situation_fallback_labelled"] = expected["situation_label_contains"] in context["situation"]["label"]
    checks[f"law_count=={expected['law_count']}"] = len(context["laws"]) == expected["law_count"]
    keys = [(law["source_file"], law["provision"]) for law in context["laws"]]
    checks["laws_deduplicated"] = len(keys) == len(set(keys))
    resolved = [source_resolves(law) for law in context["laws"]]
    checks["every_source_resolves_against_corpus"] = all(resolved)
    cited_ids = {law["chunk_id"] for law in context["laws"]}
    for chunk_id in expected["excluded_chunk_ids"]:
        # A chunk excluded in one answer may legitimately be included via another answer's
        # verification; none of these fixtures do that, so excluded means absent.
        checks[f"excluded:{chunk_id}"] = chunk_id not in cited_ids
    for fragment, label in expected["law_labels"].items():
        law = next((law for law in context["laws"] if fragment in law["title"] or fragment in law["provision"]), None)
        checks[f"label:{fragment}"] = bool(law) and law["authority"] == label
    if expected.get("notes_contain"):
        checks["note_explains_exclusion"] = any(expected["notes_contain"] in note for note in context["notes"])
    checks["uploaded_documents_listed"] = [doc["name"] for doc in context["uploaded_documents"]] == expected["uploaded_documents"]
    for secret in expected["text_absent"]:
        checks[f"redacted:{secret}"] = secret not in text
    return {
        "id": spec["id"],
        "description": spec["description"],
        "sections": context["sections"],
        "law_count": len(context["laws"]),
        "sources_resolved": f"{sum(resolved)}/{len(resolved)}",
        "checks": checks,
        "passed": all(checks.values()),
        "latency_ms": int((time.perf_counter() - started) * 1000),
    }


def evaluate_case_file(path: Path = DEFAULT_CASE_PATH) -> dict[str, Any]:
    results = [evaluate_case(spec) for spec in load_cases(path)]
    total_checks = sum(len(result["checks"]) for result in results)
    passed_checks = sum(sum(result["checks"].values()) for result in results)
    return {
        "summary": {
            "cases": len(results),
            "cases_passed": sum(result["passed"] for result in results),
            "checks": total_checks,
            "checks_passed": passed_checks,
            "llm_calls": 0,
        },
        "results": results,
    }


def render_markdown(report: dict[str, Any]) -> str:
    summary = report["summary"]
    lines = [
        "# Case File Export Evaluation",
        "",
        "## Overall Results",
        f"- Cases: {summary['cases']} ({summary['cases_passed']} passed)",
        f"- Checks: {summary['checks_passed']}/{summary['checks']} passed",
        "- LLM calls made: 0 (the case file is assembled, not generated)",
        "",
        "## Per Case",
    ]
    for result in report["results"]:
        failed = [name for name, ok in result["checks"].items() if not ok]
        lines.append(
            f"- `{result['id']}` -> {'PASS' if result['passed'] else 'FAIL'}; sections={result['sections']}; "
            f"laws={result['law_count']}; sources resolved {result['sources_resolved']}"
            + (f"; failed checks: {failed}" if failed else "")
        )
        lines.append(f"  - {result['description']}")
    lines.append("")
    return "\n".join(lines)


def write_reports(report: dict[str, Any], json_path: Path, md_path: Path) -> None:
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    md_path.write_text(render_markdown(report), encoding="utf-8")


def main(argv: list[str] | None = None) -> dict[str, Any]:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cases", type=Path, default=DEFAULT_CASE_PATH)
    parser.add_argument("--json-output", type=Path, default=DEFAULT_JSON_OUTPUT)
    parser.add_argument("--md-output", type=Path, default=DEFAULT_MD_OUTPUT)
    args = parser.parse_args(argv)
    report = evaluate_case_file(args.cases)
    write_reports(report, args.json_output, args.md_output)
    print(render_markdown(report))
    return report


if __name__ == "__main__":
    main()
