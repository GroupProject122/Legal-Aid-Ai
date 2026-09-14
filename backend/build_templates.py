"""One-time authoring script that generates the fixed docxtpl base templates under
backend/templates/. This is not a rendering path -- it is run manually (or via
`python build_templates.py`) whenever a base template needs to be created or revised, and its
output (a .docx file with Jinja2 merge tags) is what complaint_drafter.py fills in at request
time via docxtpl. Keeping this separate from complaint_drafter.py keeps the request-time module
free of python-docx paragraph/style plumbing.
"""
from __future__ import annotations

import logging

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from config import BASE_DIR

logger = logging.getLogger("legal_aid_ai.build_templates")

TEMPLATES_DIR = BASE_DIR / "templates"
TENANCY_EVICTION_TEMPLATE_PATH = TEMPLATES_DIR / "tenancy_eviction_base.docx"
CONSUMER_GOODS_TEMPLATE_PATH = TEMPLATES_DIR / "consumer_defective_goods_base.docx"
CONSUMER_ADS_TEMPLATE_PATH = TEMPLATES_DIR / "consumer_misleading_ads_base.docx"


def build_tenancy_eviction_base_template() -> None:
    """Fixed section order: heading/forum -> parties -> property & tenancy particulars ->
    grounds (selected clause inserted here) -> relief sought -> verification/signature block."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def heading(text: str, *, bold: bool = True, center: bool = True, size: int = 12) -> None:
        paragraph = document.add_paragraph()
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    def body(text: str) -> None:
        document.add_paragraph(text)

    # 1. Heading / forum
    heading("{{ forum_name }}")
    heading("{{ petition_title }}")
    document.add_paragraph()

    # 2. Parties
    heading("PARTIES", center=False)
    body("{{ landlord_name }}, residing at {{ landlord_address }}")
    body("... Petitioner/Landlord")
    heading("Versus", bold=False)
    body("{{ tenant_name }}, residing at {{ tenant_address }}")
    body("... Respondent/Tenant")
    document.add_paragraph()

    # 3. Property & tenancy particulars
    heading("PROPERTY AND TENANCY PARTICULARS", center=False)
    body("1. The tenanted premises are situated at: {{ property_address }}.")
    body("2. The agreed monthly rent of the tenanted premises is: {{ monthly_rent_display }}.")
    body("3. The tenancy commenced on: {{ tenancy_start_date_display }}.")
    document.add_paragraph()

    # 4. Grounds (selected clause inserted here)
    heading("GROUNDS FOR EVICTION", center=False)
    body("Ground relied upon: {{ grounds_label }}")
    body("{{ grounds_clause }}")
    body("{{ citation_line }}")
    document.add_paragraph()

    # 5. Relief sought
    heading("RELIEF SOUGHT", center=False)
    body("The Petitioner/Landlord respectfully prays that this Hon'ble Court be pleased to grant:")
    # docxtpl loops over whole paragraphs using {%p %} tags: one paragraph opens the loop, one
    # (or more) paragraphs are the repeated body, and one paragraph closes it. A plain
    # {% for %}...{% endfor %} inline in a single paragraph would not duplicate paragraphs.
    body("{%p for line in relief_sought_lines %}")
    body("{{ loop.index }}. {{ line }}")
    body("{%p endfor %}")
    document.add_paragraph()

    # 6. Verification / signature block
    heading("VERIFICATION", center=False)
    body(
        "I, {{ landlord_name }}, the Petitioner/Landlord above named, do hereby verify that the "
        "contents of the foregoing petition are true and correct to the best of my knowledge and "
        "belief, and that nothing material has been concealed therefrom."
    )
    body("Verified at ____________________ on {{ generated_date_display }}.")
    document.add_paragraph()
    document.add_paragraph()
    body("________________________________")
    body("{{ landlord_name }}")
    body("Petitioner/Landlord")

    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    document.save(str(TENANCY_EVICTION_TEMPLATE_PATH))
    logger.info("Saved base template to %s", TENANCY_EVICTION_TEMPLATE_PATH)


def build_consumer_defective_goods_base_template() -> None:
    """Fixed section order: heading/forum -> parties -> transaction/claim particulars ->
    grounds (selected clause inserted here) -> relief sought -> verification/signature block."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def heading(text: str, *, bold: bool = True, center: bool = True, size: int = 12) -> None:
        paragraph = document.add_paragraph()
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    def body(text: str) -> None:
        document.add_paragraph(text)

    # 1. Heading / forum
    heading("{{ forum_name }}")
    heading("{{ complaint_title }}")
    document.add_paragraph()

    # 2. Parties
    heading("PARTIES", center=False)
    body("{{ complainant_name }}, residing at {{ complainant_address }}")
    body("... Complainant")
    heading("Versus", bold=False)
    body("{{ opposite_party_name }}, having its address/place of business at {{ opposite_party_address }}")
    body("... Opposite Party")
    document.add_paragraph()

    # 3. Transaction / claim particulars
    heading("TRANSACTION PARTICULARS", center=False)
    body("1. Goods or service in question: {{ goods_or_service_description }}.")
    body("2. Date of transaction: {{ transaction_date_display }}.")
    body("3. Amount paid as consideration: {{ amount_paid_display }}.")
    body("{%p if prior_complaint_made %}")
    body("4. The Complainant previously raised this complaint with the Opposite Party, who responded as follows: {{ prior_complaint_response }}.")
    body("{%p endif %}")
    document.add_paragraph()

    # 4. Grounds (selected clause inserted here)
    heading("GROUNDS OF COMPLAINT", center=False)
    body("Ground relied upon: {{ defect_type_label }}")
    body("{{ grounds_clause }}")
    body("{{ citation_line }}")
    body("{{ forum_citation_line }}")
    document.add_paragraph()

    # 5. Relief sought
    heading("RELIEF SOUGHT", center=False)
    body("The Complainant respectfully prays that this Hon'ble Commission be pleased to direct the Opposite Party to:")
    body("{%p for line in relief_sought_lines %}")
    body("{{ loop.index }}. {{ line }}")
    body("{%p endfor %}")
    document.add_paragraph()

    # 6. Verification / signature block
    heading("VERIFICATION", center=False)
    body(
        "I, {{ complainant_name }}, the Complainant above named, do hereby verify that the contents "
        "of the foregoing complaint are true and correct to the best of my knowledge and belief, and "
        "that nothing material has been concealed therefrom."
    )
    body("Verified at ____________________ on {{ generated_date_display }}.")
    document.add_paragraph()
    document.add_paragraph()
    body("________________________________")
    body("{{ complainant_name }}")
    body("Complainant")

    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    document.save(str(CONSUMER_GOODS_TEMPLATE_PATH))
    logger.info("Saved base template to %s", CONSUMER_GOODS_TEMPLATE_PATH)


def build_misleading_advertisement_base_template() -> None:
    """Fixed section order: heading/forum -> parties -> transaction/claim particulars ->
    grounds (selected clause inserted here) -> relief sought -> verification/signature block."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def heading(text: str, *, bold: bool = True, center: bool = True, size: int = 12) -> None:
        paragraph = document.add_paragraph()
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    def body(text: str) -> None:
        document.add_paragraph(text)

    # 1. Heading / forum
    heading("{{ forum_name }}")
    heading("{{ complaint_title }}")
    document.add_paragraph()

    # 2. Parties
    heading("PARTIES", center=False)
    body("{{ complainant_name }}, residing at {{ complainant_address }}")
    body("... Complainant")
    heading("Versus", bold=False)
    body("{{ opposite_party_name }}, having its address/place of business at {{ opposite_party_address }}")
    body("... Opposite Party")
    document.add_paragraph()

    # 3. Transaction / claim particulars
    heading("ADVERTISEMENT/PRACTICE PARTICULARS", center=False)
    body("1. Advertisement or practice complained of: {{ advertisement_or_practice_description }}.")
    body("2. Platform or medium: {{ platform_label }}.")
    body("3. Date encountered: {{ date_encountered_display }}.")
    # amount_paid is optional in this scenario (not every claim -- e.g. a dark pattern
    # encountered without completing a purchase -- involves a completed transaction), so this
    # line is only shown when it was actually supplied. Items 1-3 above are always present, so
    # hardcoding "4." here (rather than a loop index) is safe -- there's no gap to renumber around.
    body("{%p if amount_paid_display %}")
    body("4. Amount paid: {{ amount_paid_display }}.")
    body("{%p endif %}")
    document.add_paragraph()

    # 4. Grounds (selected clause inserted here)
    heading("GROUNDS OF COMPLAINT", center=False)
    body("Ground relied upon: {{ claim_type_label }}")
    body("{{ grounds_clause }}")
    body("{{ citation_line }}")
    document.add_paragraph()

    # 5. Relief sought
    heading("RELIEF SOUGHT", center=False)
    body("The Complainant respectfully prays that this Hon'ble Commission be pleased to direct the Opposite Party to:")
    body("{%p for line in relief_sought_lines %}")
    body("{{ loop.index }}. {{ line }}")
    body("{%p endfor %}")
    document.add_paragraph()

    # 6. Verification / signature block
    heading("VERIFICATION", center=False)
    body(
        "I, {{ complainant_name }}, the Complainant above named, do hereby verify that the contents "
        "of the foregoing complaint are true and correct to the best of my knowledge and belief, and "
        "that nothing material has been concealed therefrom."
    )
    body("Verified at ____________________ on {{ generated_date_display }}.")
    document.add_paragraph()
    document.add_paragraph()
    body("________________________________")
    body("{{ complainant_name }}")
    body("Complainant")

    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    document.save(str(CONSUMER_ADS_TEMPLATE_PATH))
    logger.info("Saved base template to %s", CONSUMER_ADS_TEMPLATE_PATH)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    build_tenancy_eviction_base_template()
    build_consumer_defective_goods_base_template()
    build_misleading_advertisement_base_template()


if __name__ == "__main__":
    main()
