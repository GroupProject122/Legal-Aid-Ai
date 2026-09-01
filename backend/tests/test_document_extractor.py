from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import document_extractor
import main
from config import BASE_DIR, DOCUMENTS_DIR, INDEX_PATH, METADATA_PATH


client = TestClient(main.app)


def pdf_bytes(*pages: str) -> bytes:
    objects: list[bytes] = [
        f"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n".encode("latin-1"),
        f"2 0 obj << /Type /Pages /Kids [{' '.join(f'{3 + index * 3} 0 R' for index in range(len(pages)))}] /Count {len(pages)} >> endobj\n".encode("latin-1"),
    ]
    for index, text in enumerate(pages):
        page_id = 3 + index * 3
        font_id = page_id + 1
        content_id = page_id + 2
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.extend(
            [
                f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >> endobj\n".encode("latin-1"),
                f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode("latin-1"),
                f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode("latin-1") + stream + b"\nendstream endobj\n",
            ]
        )
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for obj in objects:
        offsets.append(len(output))
        output.extend(obj)
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1"))
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("latin-1"))
    return bytes(output)


def blank_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Rent agreement between landlord and tenant.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Monthly rent"
    table.cell(0, 1).text = "10000"
    document.save(buffer)
    return buffer.getvalue()


def test_valid_pdf_text_extraction_and_page_count():
    result = document_extractor.extract_document("receipt.pdf", pdf_bytes("Invoice page one", "Payment page two"), "application/pdf")

    assert result.status == "success"
    assert result.file_type == "pdf"
    assert result.page_count == 2
    assert "Invoice page one" in result.text
    assert result.pages[1].page_number == 2


def test_valid_docx_extracts_paragraph_and_table_text():
    result = document_extractor.extract_document(
        "rent_agreement.docx",
        docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.status == "success"
    assert result.file_type == "docx"
    assert "Rent agreement" in result.text
    assert "Monthly rent | 10000" in result.text


def test_valid_txt_extraction():
    result = document_extractor.extract_document("notice.txt", b"Legal notice text\n\nSecond paragraph.", "text/plain")

    assert result.status == "success"
    assert result.file_type == "txt"
    assert "Second paragraph" in result.text


def test_unsupported_extension_rejected():
    result = document_extractor.extract_document("archive.zip", b"PK\x03\x04binary", "application/zip")

    assert result.status == "unsupported"
    assert result.text == ""


def test_oversized_file_rejected():
    content = b"x" * (document_extractor.max_upload_bytes() + 1)

    with pytest.raises(document_extractor.DocumentExtractionError):
        document_extractor.extract_document("large.txt", content, "text/plain")


def test_empty_text_pdf_returns_ocr_required():
    result = document_extractor.extract_document("scan.pdf", blank_pdf_bytes(), "application/pdf")

    assert result.status == "ocr_required"
    assert result.page_count == 1
    assert any("scanned" in warning for warning in result.warnings)


def test_binary_like_txt_rejected():
    result = document_extractor.extract_document("data.txt", b"\x00\x01\x02\x03not text", "text/plain")

    assert result.status == "failed"
    assert "binary" in " ".join(result.warnings)


def test_endpoint_returns_expected_schema():
    response = client.post(
        "/api/documents/extract",
        files={"file": ("invoice.txt", b"Consumer invoice for phone purchase.", "text/plain")},
    )

    assert response.status_code == 200
    data = response.json()
    assert set(data) >= {
        "status",
        "filename",
        "file_type",
        "size_bytes",
        "page_count",
        "character_count",
        "text",
        "pages",
        "warnings",
    }
    assert data["status"] in {"success", "partial"}


def test_endpoint_errors_do_not_expose_raw_document_text(monkeypatch):
    secret_text = "Aadhaar 1234 private bank text"

    def fail_extract(*_args, **_kwargs):
        raise RuntimeError(secret_text)

    monkeypatch.setattr(main.document_extractor, "extract_document", fail_extract)
    response = client.post(
        "/api/documents/extract",
        files={"file": ("notice.txt", secret_text.encode("utf-8"), "text/plain")},
    )

    assert response.status_code == 500
    assert secret_text not in response.text


def test_upload_endpoint_does_not_touch_legal_corpus_or_vectorstore():
    tracked_paths = [DOCUMENTS_DIR / "corpus_manifest.json", BASE_DIR / "parsed" / "legal_chunks.jsonl", INDEX_PATH, METADATA_PATH]
    before = {path: path.stat().st_mtime_ns for path in tracked_paths if path.exists()}

    response = client.post(
        "/api/documents/extract",
        files={"file": ("receipt.txt", b"Payment receipt for online order.", "text/plain")},
    )

    assert response.status_code == 200
    after = {path: path.stat().st_mtime_ns for path in tracked_paths if path.exists()}
    assert after == before
