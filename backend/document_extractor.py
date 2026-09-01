from __future__ import annotations

import io
import logging
import re
import time
from dataclasses import asdict, dataclass, field
from typing import Any

from docx import Document
from pypdf import PdfReader

from config import MAX_UPLOAD_SIZE_MB

logger = logging.getLogger("legal_aid_ai.document_extractor")

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}
SUPPORTED_MIME_HINTS = {
    "pdf": {"application/pdf"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
    "txt": {"text/plain", "application/octet-stream"},
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"}
MAX_PREVIEW_CHARS = 5000


class DocumentExtractionError(ValueError):
    pass


@dataclass
class ExtractedPage:
    page_number: int
    text: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class ExtractionResult:
    status: str
    filename: str
    file_type: str | None
    size_bytes: int
    page_count: int | None = None
    character_count: int = 0
    text: str = ""
    pages: list[ExtractedPage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    processing_time_ms: int | None = None

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["pages"] = [page.to_dict() for page in self.pages]
        return data


def extract_document(filename: str, content: bytes, content_type: str | None = None) -> ExtractionResult:
    started = time.perf_counter()
    clean_name = safe_filename(filename)
    size = len(content)
    file_type = detect_file_type(clean_name, content, content_type)
    if size > max_upload_bytes():
        raise DocumentExtractionError(f"File is too large. Maximum supported size is {MAX_UPLOAD_SIZE_MB} MB.")
    if file_type is None:
        return result(
            "unsupported",
            clean_name,
            None,
            size,
            warnings=["This file type is not supported yet. Please upload PDF, DOCX, or TXT."],
            started=started,
        )
    if not content:
        return result("failed", clean_name, file_type, size, warnings=["The uploaded file is empty."], started=started)
    try:
        if file_type == "pdf":
            extracted = extract_pdf(clean_name, content, size)
        elif file_type == "docx":
            extracted = extract_docx(clean_name, content, size)
        elif file_type == "txt":
            extracted = extract_txt(clean_name, content, size)
        else:
            extracted = result("unsupported", clean_name, file_type, size, warnings=["Unsupported file type."], started=started)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        logger.warning("Document extraction failed for type=%s size=%s: %s", file_type, size, exc.__class__.__name__)
        extracted = result("failed", clean_name, file_type, size, warnings=["Text extraction failed for this document."], started=started)
    extracted.processing_time_ms = int((time.perf_counter() - started) * 1000)
    return extracted


def detect_file_type(filename: str, content: bytes, content_type: str | None = None) -> str | None:
    extension = extension_of(filename)
    if extension in IMAGE_EXTENSIONS:
        return None
    hinted_type = SUPPORTED_EXTENSIONS.get(extension)
    if content.startswith(b"%PDF"):
        return "pdf" if hinted_type in {None, "pdf"} else None
    if content.startswith(b"PK\x03\x04"):
        return "docx" if hinted_type == "docx" else None
    if hinted_type == "txt" and looks_like_text(content):
        return "txt"
    if hinted_type and content_type:
        normalized_content_type = content_type.split(";")[0].strip().lower()
        if normalized_content_type in SUPPORTED_MIME_HINTS.get(hinted_type, set()):
            return hinted_type
    if hinted_type == "txt" and not is_binary_like(content):
        return "txt"
    return None


def extract_pdf(filename: str, content: bytes, size: int) -> ExtractionResult:
    reader = PdfReader(io.BytesIO(content))
    pages: list[ExtractedPage] = []
    blank_pages = 0
    for index, page in enumerate(reader.pages, start=1):
        page_text = normalize_text(page.extract_text() or "")
        if not page_text:
            blank_pages += 1
        pages.append(ExtractedPage(page_number=index, text=page_text))
    text = join_page_text(pages)
    warnings = quality_warnings(text, blank_pages, len(pages))
    if len(reader.pages) and is_ocr_required(text, len(reader.pages), blank_pages):
        return ExtractionResult(
            status="ocr_required",
            filename=filename,
            file_type="pdf",
            size_bytes=size,
            page_count=len(reader.pages),
            character_count=len(text),
            text=text,
            pages=pages,
            warnings=[
                "This document appears to be scanned or image-based, so reliable text extraction is not available in the current phase."
            ],
        )
    status = "success" if text else "failed"
    if status == "failed":
        warnings.append("No readable text could be extracted from this PDF.")
    elif warnings:
        status = "partial"
    return ExtractionResult(status, filename, "pdf", size, len(reader.pages), len(text), text, pages, warnings)


def extract_docx(filename: str, content: bytes, size: int) -> ExtractionResult:
    document = Document(io.BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                blocks.append(line)
    text = normalize_text("\n\n".join(blocks))
    warnings = quality_warnings(text, 0, None)
    status = "success" if text else "failed"
    if status == "failed":
        warnings.append("No readable text could be extracted from this DOCX file.")
    elif warnings:
        status = "partial"
    return ExtractionResult(status, filename, "docx", size, None, len(text), text, [], warnings)


def extract_txt(filename: str, content: bytes, size: int) -> ExtractionResult:
    if is_binary_like(content):
        return ExtractionResult(
            status="failed",
            filename=filename,
            file_type="txt",
            size_bytes=size,
            character_count=0,
            warnings=["This TXT file appears to contain binary or unreadable content."],
        )
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("TXT files must be readable UTF-8 text.") from exc
    text = normalize_text(text)
    warnings = quality_warnings(text, 0, None)
    status = "success" if text else "failed"
    if status == "failed":
        warnings.append("No readable text could be extracted from this TXT file.")
    elif warnings:
        status = "partial"
    return ExtractionResult(status, filename, "txt", size, None, len(text), text, [], warnings)


def result(status: str, filename: str, file_type: str | None, size: int, warnings: list[str], started: float) -> ExtractionResult:
    return ExtractionResult(
        status=status,
        filename=filename,
        file_type=file_type,
        size_bytes=size,
        warnings=warnings,
        processing_time_ms=int((time.perf_counter() - started) * 1000),
    )


def normalize_text(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def join_page_text(pages: list[ExtractedPage]) -> str:
    return normalize_text("\n\n".join(page.text for page in pages if page.text))


def quality_warnings(text: str, blank_pages: int, page_count: int | None) -> list[str]:
    warnings: list[str] = []
    if text and len(text) < 20:
        warnings.append("Very little readable text was extracted.")
    if page_count and blank_pages / page_count >= 0.5:
        warnings.append("Many PDF pages had no extractable text.")
    return warnings


def is_ocr_required(text: str, page_count: int, blank_pages: int) -> bool:
    if page_count <= 0:
        return False
    return not text or len(text) < page_count * 5 or blank_pages == page_count


def safe_filename(filename: str) -> str:
    return (filename or "uploaded_document").split("/")[-1].split("\\")[-1]


def extension_of(filename: str) -> str:
    match = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return match.group(1).lower() if match else ""


def max_upload_bytes() -> int:
    return MAX_UPLOAD_SIZE_MB * 1024 * 1024


def looks_like_text(content: bytes) -> bool:
    if not content:
        return True
    sample = content[:2048]
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return not is_binary_like(sample)


def is_binary_like(content: bytes) -> bool:
    if not content:
        return False
    sample = content[:4096]
    if b"\x00" in sample:
        return True
    control = sum(1 for byte in sample if byte < 32 and byte not in {9, 10, 12, 13})
    return control / max(len(sample), 1) > 0.08
